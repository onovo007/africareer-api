import io,json,unittest
from datetime import date
from unittest.mock import patch,Mock
from fastapi.testclient import TestClient
from docx import Document
import api,core,evidence,feedback_store
from quality import DraftValidationError,opportunity_matches
from applications import application_rules,check_sections,render_application
from draft_review import audited_draft

class RemediationTests(unittest.TestCase):
 def setUp(self):
  api.limiter.enabled=False
  self.client=TestClient(api.app)
 def test_evidence_has_no_accuracy_probability(self):
  status=evidence.evidence_status([])
  self.assertIsNone(status['confidence_score'])
  self.assertEqual(status['status'],'general_guidance')
  self.assertEqual(status['claim_support'],'Not independently verified')
 def test_reference_integrity_scope_and_expiry(self):
  refs=evidence.references(date(2026,9,28))
  self.assertGreaterEqual(len(refs),9)
  self.assertTrue(all(r['active'] for r in refs))
  self.assertTrue(all(not r['active'] for r in evidence.references(date(2026,11,1))))
  self.assertNotIn('cambridge-phpc',[r['id'] for r in evidence.select_references('PhD at University of Ghana')])
 def test_knowledge_endpoint_is_public_without_credentials(self):
  r=self.client.get('/knowledge');self.assertEqual(r.status_code,200)
  self.assertNotIn('api_key',r.text.lower())
 def test_no_pinecone_still_has_inspectable_local_references(self):
  with patch.object(core,'PINECONE_API_KEY',None):
   ctx,sources=core._retrieve('World Bank Africa jobs digital')
  self.assertIn('World Bank',ctx)
  self.assertTrue(any('worldbank.org' in s for s in sources))
 def test_assistant_without_search_still_answers(self):
  with patch.object(core,'TAVILY_API_KEY',''),patch.object(core,'safe_llm_call',return_value='Make a plan'):
   r=self.client.post('/assistant',json={'question':'How do I change careers?'})
  self.assertEqual(r.status_code,200);self.assertIn('evidence',r.json())
 def test_do_not_send_notebook_to_web_search(self):
  with patch.object(core,'TAVILY_API_KEY','test'),patch.object(core,'web_search_links',return_value=[]) as search,patch.object(core,'safe_llm_call',return_value='Advice'):
   core.assistant_answer('USER-EDITED NOTEBOOK\nPrivate fictional name\nCURRENT USER MESSAGE\nFind scholarships')
  self.assertEqual(search.call_args.args[0],'Find scholarships')
 def test_factual_review_catches_word_claim_and_repairs(self):
  values=iter([json.dumps({'professional_summary':'Award-winning analyst'}),json.dumps({'passed':False,'issues':['No award supplied']}),json.dumps({'professional_summary':'Analyst'}),json.dumps({'passed':True,'issues':[]})])
  self.assertEqual(audited_draft('CV','Analyst',lambda *a:next(values),json.loads)['professional_summary'],'Analyst')
 def test_malformed_or_failed_audit_fails_closed(self):
  with self.assertRaises(DraftValidationError):
   audited_draft('CV','Analyst',lambda *a:'{}',json.loads)
 def test_ucas_preset_and_boundaries(self):
  rules=application_rules('Undergraduate program','Oxford','Computer Science')
  sections=[{'text':'a'*350}]*3
  self.assertEqual(check_sections(sections,rules)[1]['characters'],1050)
  for bad in ([{'text':'a'*349}]*3,[{'text':'a'*1500}]*3):
   with self.assertRaises(DraftValidationError):check_sections(bad,rules)
 def test_cambridge_limits_are_not_global(self):
  rules=application_rules('PhD / Doctorate position','Cambridge','Public Health and Primary Care')
  self.assertEqual(rules['max_characters'],2500)
  self.assertIsNone(application_rules('PhD / Doctorate position','Cambridge','History')['max_characters'])
  with self.assertRaises(DraftValidationError):check_sections([{'text':'x'*2501}],rules)
 def test_user_limit_cannot_loosen_preset(self):
  rules=application_rules('Undergraduate program','Oxford','Maths','ucas',max_characters=10000)
  self.assertEqual(rules['max_characters'],4000)
 def test_statement_docx_has_only_statement_no_address_or_signature(self):
  rules=application_rules('PhD / Doctorate position','Cambridge','Public Health and Primary Care')
  raw=render_application({'rules':rules,'sections':[{'text':'I studied epidemiology.\n\nI hope to study prevention.'}]})
  doc=Document(io.BytesIO(raw));self.assertEqual([p.text for p in doc.paragraphs],['I studied epidemiology.','I hope to study prevention.'])
  self.assertEqual(len(doc.tables),0)
 def test_export_checks_user_confirmation_and_counts(self):
  body={'category':'Undergraduate program','school':'Oxford','programme':'Maths','background':'I studied maths','sections':[{'text':'x'*1400}]*3,'confirmed':True}
  self.assertEqual(self.client.post('/application-document',json=body).status_code,409)
  body['confirmed']=False
  self.assertEqual(self.client.post('/application-document',json=body).status_code,422)
 def test_europe_phd_rejects_us_masters_and_lecturer(self):
  for title in ('Masters in public health, USA','Lecturer in public health PhD, Germany','PhD public health United States'):
   self.assertFalse(opportunity_matches({'title':title,'url':'https://university.example/course'},'PhD / Doctorate','public health','Europe'))
  self.assertTrue(opportunity_matches({'title':'PhD in public health','content':'Doctoral study in Germany','url':'https://university.example/course'},'PhD / Doctorate','public health','Europe'))
 def test_feedback_retries_use_same_id_and_readback(self):
  payload={'request_id':'f7254255-9223-420e-8c8e-98539d116053','rating':'down'}
  client=Mock();client.post.return_value.status_code=201;client.get.return_value.status_code=200;client.get.return_value.json.return_value=[payload]
  with patch.object(feedback_store.httpx,'Client') as factory:
   factory.return_value.__enter__.return_value=client
   self.assertTrue(feedback_store.save_feedback('https://project.example','server-key',payload))
   self.assertTrue(feedback_store.save_feedback('https://project.example','server-key',payload))
  self.assertEqual(client.post.call_args.kwargs['params'],{'on_conflict':'request_id'})
  client.get.return_value.json.return_value=[dict(payload,rating='up')]
  with patch.object(feedback_store.httpx,'Client') as factory:
   factory.return_value.__enter__.return_value=client
   self.assertFalse(feedback_store.save_feedback('https://project.example','server-key',payload))
 def test_pdf_empty_scan_and_table_docx(self):
  from pypdf import PdfWriter
  pdf=PdfWriter();pdf.add_blank_page(width=612,height=792);out=io.BytesIO();pdf.write(out)
  self.assertEqual(self.client.post('/extract-text',files={'file':('scan.pdf',out.getvalue())}).status_code,400)
  self.assertEqual(self.client.post('/extract-text',files={'file':('broken.pdf',b'not a pdf')}).status_code,400)


 def test_table_resume_extraction_keeps_qualifications(self):
  doc=Document();table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Education';table.cell(0,1).text='BSc Statistics, expected June 2027'
  out=io.BytesIO();doc.save(out)
  result=self.client.post('/extract-text',files={'file':('table.docx',out.getvalue())})
  self.assertEqual(result.status_code,200);self.assertIn('expected June 2027',result.json()['text'])
 def test_job_page_metadata_dates_location_and_missing_fields(self):
  from job_verification import checked_posting
  from datetime import datetime,timezone
  now=datetime(2026,9,28,tzinfo=timezone.utc)
  job={'@type':'JobPosting','title':'Data analyst','hiringOrganization':{'name':'Example NGO'},'jobLocation':{'address':{'addressCountry':'Kenya'}},'jobLocationType':'TELECOMMUTE','datePosted':'2026-09-25','validThrough':'2026-10-01','description':'Mid level public health data analyst'}
  def check(j,location='Kenya'):
   return checked_posting('<script type="application/ld+json">'+json.dumps(j)+'</script>','https://example.org/careers/123','Data analyst',location,'Mid level','Remote','public health','Past week',now)
  self.assertIsNotNone(check(job))
  self.assertIsNone(check(job,'Nigeria'))
  self.assertIsNone(check(dict(job,validThrough='2026-09-20')))
  self.assertIsNone(check(dict(job,datePosted=None)))
  self.assertIsNone(check(dict(job,jobLocationType='ON_SITE')))


 def test_legacy_course_price_alias_never_invokes_model(self):
  with patch.object(core,'verify_url',return_value=True),patch.object(core,'safe_llm_call') as model:
   self.assertTrue(core.find_courses('Python','Beginner','Free'))
   self.assertTrue(core.find_courses('Excel','Intermediate','Paid'))
   model.assert_not_called()
 def test_production_cors_is_explicit_and_untrusted_origin_denied(self):
  headers={'Origin':'https://africareer.quantiuminsights.com','Access-Control-Request-Method':'POST'}
  self.assertEqual(self.client.options('/assistant',headers=headers).status_code,200)
  headers['Origin']='https://unknown.example'
  self.assertEqual(self.client.options('/assistant',headers=headers).status_code,400)

 def test_phd_search_excludes_traineeship_with_doctoral_eligibility(self):
  self.assertFalse(opportunity_matches({'title':'ECDC Traineeship Programme 2027', 'content':'Public health graduates with a PhD may apply in Europe', 'url':'https://example.org/apply'},'PhD / Doctorate','public health','Europe'))

if __name__=='__main__':unittest.main()
